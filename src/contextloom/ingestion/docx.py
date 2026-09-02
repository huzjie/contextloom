"""DOCX 解析器（可选依赖 python-docx）。"""

from __future__ import annotations

from pathlib import Path

from contextloom.ingestion.base import BaseParser, ParsedDocument
from contextloom.ingestion.registry import register_parser


@register_parser(".docx")
class DocxParser(BaseParser):
    extensions = (".docx",)

    async def parse(self, path: str | Path) -> ParsedDocument:
        p = Path(path)
        try:
            import docx
        except ImportError as exc:
            raise RuntimeError("解析 DOCX 需要安装：pip install 'contextloom[docx]'") from exc
        d = docx.Document(str(p))
        paragraphs = [para.text for para in d.paragraphs if para.text.strip()]
        headings = [para.text for para in d.paragraphs
                    if para.style.name.lower().startswith("heading")]
        text = "\n\n".join(paragraphs)
        return ParsedDocument(title=headings[0] if headings else p.stem, text=text,
                              headings=headings, source=str(p),
                              metadata={"format": "docx", "paragraphs": len(paragraphs)})
